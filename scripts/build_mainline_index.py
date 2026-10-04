"""Build the runtime mainline index; not used by the game at runtime.

Usage: python scripts/build_mainline_index.py --quest-title-lang path/to/QuestTitle.lang
The workbook supplies world/order/title; QuestTitle.lang supplies verified
language keys and Chinese title variants. Neither source contains Quest IDs.
Unchanged world/title rows retain the existing index's verified identities.
"""

import argparse
import json
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "W101 Mainline Quests.xlsx"
OUTPUT = ROOT / "src" / "data" / "mainline_quests.json"


def title_key(value):
    value = unicodedata.normalize("NFKC", value).casefold()
    return "".join(re.findall(r"[a-z0-9\u3400-\u9fff]+", value))


def split_title(value):
    text = unicodedata.normalize("NFKC", str(value)).strip()
    match = re.match(r"^(.*?)\s+\((.+)\)$", text)
    if match:
        return match[1].strip(), match[2].strip()
    return text, ""


def workbook_rows(workbook_path=WORKBOOK):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        cells = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    headers = cells[1]
    worlds = []
    for column, header in enumerate(headers):
        match = re.fullmatch(r"\s*(.+?)\s*\((\d+)\)\s*", str(header or ""))
        if match:
            worlds.append((column, match[1].strip(), int(match[2])))
    rows = []
    for cells in cells[2:]:
        for column, world, total in worlds:
            match = re.match(r"^\s*(\d+)\.\s*(.+)$", str(cells[column] or ""))
            if not match:
                continue
            english, note = split_title(match[2])
            aliases = []
            old_name = re.search(r"\b(?:formerly|previously|old name)\s*:?[\s\u00a0]+(.+)", note, re.I)
            if old_name:
                aliases.append(old_name[1].strip(' \"“”'))
            rows.append(dict(world=world, number=int(match[1]), total=total,
                             english=english, quest_ids=[], keys=[], chinese=[],
                             aliases=aliases, note=note))
    mismatches = []
    for _, world, total in worlds:
        numbers = sorted(row["number"] for row in rows if row["world"] == world)
        if numbers != list(range(1, len(numbers) + 1)):
            raise ValueError(f"{world}: numbering is not consecutive")
        if len(numbers) != total:
            mismatches.append({"world": world, "header_total": total,
                               "numbered_rows": len(numbers)})
        for row in rows:
            if row["world"] == world:
                row["total"] = len(numbers)
    return rows, mismatches


def language_titles(path):
    lines = Path(path).read_bytes().decode("utf-16").splitlines()
    if not lines or not lines[0].endswith(":QuestTitle"):
        raise ValueError("Expected a QuestTitle.lang language table")
    by_title = defaultdict(list)
    for index in range(1, len(lines) - 2, 3):
        code, english, chinese = lines[index:index + 3]
        if code and english:
            by_title[title_key(english)].append((f"QuestTitle_{code}", chinese.strip()))
    return by_title


def build(lang_path, workbook_path=WORKBOOK, output=OUTPUT, previous_index=OUTPUT,
          title_source='QuestTitle.lang bilingual snapshot'):
    titles = language_titles(lang_path)
    rows, mismatches = workbook_rows(workbook_path)
    previous = defaultdict(list)
    if previous_index is not None and previous_index.exists():
        for row in json.loads(previous_index.read_text(encoding='utf-8'))['rows']:
            previous[(title_key(row['world']), title_key(row['english']))].append(row)
    for row in rows:
        old_rows = previous[(title_key(row['world']), title_key(row['english']))]
        if len(old_rows) > 1:
            raise ValueError(f"{row['world']}: ambiguous existing title {row['english']!r}")
        if old_rows:
            old = old_rows[0]
            for field in ('quest_ids', 'keys', 'chinese'):
                row[field] = list(old.get(field, []))
            row['aliases'] = list(dict.fromkeys([*row['aliases'], *old.get('aliases', [])]))
            if old.get('quest_ids') and old.get('note'):
                row['note'] = old['note']
            continue
        found = []
        for title in [row["english"], *row["aliases"]]:
            found.extend(titles.get(title_key(title), []))
        row["keys"] = list(dict.fromkeys(code for code, _ in found))
        row["chinese"] = list(dict.fromkeys(
            translation for _, translation in found
            if re.search(r"[\u3400-\u9fff]", translation)
        ))
    data = {
        "source": workbook_path.name,
        "title_source": title_source,
        "quest_id_source": "Neither source contains Quest IDs; verified IDs may be added later",
        "source_total_mismatches": mismatches,
        "rows": rows,
    }
    output.write_text(
        "{\n" + ",\n".join(
            f'  "{key}": {json.dumps(value, ensure_ascii=False)}'
            for key, value in data.items() if key != "rows"
        ) + ',\n  "rows": [\n' +
        ",\n".join("    " + json.dumps(row, ensure_ascii=False, separators=(",", ":")) for row in rows) +
        "\n  ]\n}\n", encoding="utf-8",
    )
    print(f"{len(rows)} rows; {sum(bool(r['keys']) for r in rows)} with language keys; "
          f"{sum(bool(r['chinese']) for r in rows)} with Chinese titles; output {output}")
    return data


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--quest-title-lang", required=True, type=Path)
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--previous-index", type=Path, default=OUTPUT)
    parser.add_argument("--title-source", default='QuestTitle.lang bilingual snapshot')
    args = parser.parse_args()
    build(args.quest_title_lang, args.workbook, args.output, args.previous_index, args.title_source)
